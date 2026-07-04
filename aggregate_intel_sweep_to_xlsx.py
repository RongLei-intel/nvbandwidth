#!/usr/bin/env python3
from __future__ import annotations

import argparse
import re
import os
import glob
from pathlib import Path

import openpyxl

DATA_SOURCE = "intel_sweep_20260704_ro_sweep"
SYSTEM = "intel"
SOURCE_KIND = "intel_point_summary"

CASE_NAME_RE = re.compile(
    r"^t(?P<testcase>\d+)_(?P<unit>kib|mib)_b(?P<buf_num>\d+)(?P<buf_unit>KiB|MiB)_"
    r"(?P<param_prefix>tsmCopyBytes|tptrChaseLoadBytes)(?P<param_val>\d+)$"
)

COLUMNS = [
    "data_source", "system", "testcase", "run_label", "ro_state", "config_tag",
    "tuning_param_name", "per_load_bytes", "buffer_MiB", "repeat_count", "statuses",
    "source_kind", "source_file", "source_note", "loop_count", "nvbandwidth_gbps",
    "memory_read_gbps", "memory_gbps", "pcie_read_gbps", "pcie_gbps", "gpu_pcie_gbps",
    "target_socket_mem_read_gbps", "pcm_memory_samples", "pcm_pcie_samples",
    "nvidia_dmon_rows", "run_seconds",
]


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Aggregate intel sweep summary.txt files into xlsx")
    p.add_argument("--root", type=Path, default=Path(__file__).resolve().parent / "out")
    p.add_argument("--xlsx", type=Path, default=Path(__file__).resolve().parent / "all_point_summary_compact.xlsx")
    p.add_argument("--data-source", type=str, default=DATA_SOURCE)
    p.add_argument("--dry-run", action="store_true")
    return p.parse_args()


def parse_backtick_value(text: str, key: str) -> str:
    m = re.search(rf"{re.escape(key)}:\s*`([^`]*)`", text)
    return m.group(1) if m else ""


def parse_table_metric(content: str, section_pattern: str, metric_name: str) -> str:
    """Extract the 'avg' column value for a metric in a markdown table within a section."""
    section_idx = content.find(section_pattern)
    if section_idx == -1:
        return ""
    section = content[section_idx:]

    escaped = re.escape(metric_name)
    m = re.search(rf"^\|\s*{escaped}\s*\|\s*([^\s|]+)", section, re.MULTILINE)
    if m:
        return m.group(1)
    return ""


def parse_row_count(content: str, section_pattern: str, label: str) -> str:
    section_idx = content.find(section_pattern)
    if section_idx == -1:
        return ""
    section = content[section_idx:]
    m = re.search(rf"{re.escape(label)}:\s*`(\d+)`", section)
    return m.group(1) if m else ""


def parse_case_name(dirname: str) -> dict[str, str]:
    m = CASE_NAME_RE.match(dirname)
    if not m:
        return {}
    d = m.groupdict()
    buf_val = float(d["buf_num"])
    buf_unit = d["buf_unit"]
    buffer_mib = buf_val if buf_unit == "MiB" else buf_val / 1024.0

    param_prefix = d["param_prefix"]
    if param_prefix == "tsmCopyBytes":
        tuning_name = "sm_copy_bytes"
    elif param_prefix == "tptrChaseLoadBytes":
        tuning_name = "ptr_chase_load_bytes"
    else:
        tuning_name = param_prefix

    testcase_num = d["testcase"]
    testcase_label = "Sequential" if testcase_num == "16" else "Random"

    return {
        "testcase": testcase_label,
        "tuning_param_name": tuning_name,
        "per_load_bytes": d["param_val"],
        "buffer_MiB": str(buffer_mib),
    }


def parse_run_label(dirpath: str) -> dict[str, str]:
    parts = Path(dirpath).parts
    run_label = parts[-2] if len(parts) >= 2 else ""
    if run_label.startswith("ro_on_"):
        ro_state = "on"
        config_tag = run_label[len("ro_on_"):]
    elif run_label.startswith("ro_off_"):
        ro_state = "off"
        config_tag = run_label[len("ro_off_"):]
    else:
        ro_state = ""
        config_tag = ""
    return {"run_label": run_label, "ro_state": ro_state, "config_tag": config_tag}


def parse_summary_file(filepath: str, root_dir: str) -> dict[str, str] | None:
    content = Path(filepath).read_text(errors="replace")
    case_dir = str(Path(filepath).parent)
    case_name = os.path.basename(case_dir)

    case_info = parse_case_name(case_name)
    run_info = parse_run_label(case_dir)

    if not case_info or not run_info.get("run_label"):
        print(f"  SKIP: cannot parse {filepath}")
        return None

    loop_count = parse_backtick_value(content, "loop_count")
    run_seconds = parse_backtick_value(content, "run_seconds")
    nvbandwidth_rc = parse_backtick_value(content, "nvbandwidth_rc")

    nvbandwidth_gbps = parse_table_metric(content, "## nvbandwidth", "host_to_device_memcpy_sm")
    if not nvbandwidth_gbps:
        nvbandwidth_gbps = parse_table_metric(content, "## nvbandwidth", "host_device_bandwidth_sm")

    memory_read_gbps = parse_table_metric(content, "## PCM memory bandwidth", "System.Read")
    memory_gbps = parse_table_metric(content, "## PCM memory bandwidth", "System.Memory")
    pcie_read_gbps = parse_table_metric(content, "## PCM PCIe bandwidth", "System PCIe Rd (B)")
    pcie_gbps = parse_table_metric(content, "## PCM PCIe bandwidth", "System PCIe Total (B)")
    gpu_pcie_gbps = parse_table_metric(content, "## NVIDIA dmon GPU telemetry", "All GPUs rxpci+txpci")
    target_socket_mem_read_gbps = parse_table_metric(
        content, "## PCM memory bandwidth", "SKT0.Mem Read (MB/s)"
    )

    pcm_mem_samples = parse_row_count(content, "## PCM memory bandwidth", "Samples")
    pcm_pcie_samples = parse_row_count(content, "## PCM PCIe bandwidth", "Samples")
    nvidia_dmon_rows = parse_row_count(content, "## NVIDIA dmon GPU telemetry", "Rows")

    statuses = "ok" if nvbandwidth_rc == "0" else f"rc={nvbandwidth_rc}"

    abs_path = str(Path(filepath).resolve())

    return {
        "data_source": DATA_SOURCE,
        "system": SYSTEM,
        **case_info,
        **run_info,
        "repeat_count": "1",
        "statuses": statuses,
        "source_kind": SOURCE_KIND,
        "source_file": abs_path,
        "source_note": f"GNR RO sweep results {run_info['run_label']}",
        "loop_count": loop_count,
        "nvbandwidth_gbps": nvbandwidth_gbps,
        "memory_read_gbps": memory_read_gbps,
        "memory_gbps": memory_gbps,
        "pcie_read_gbps": pcie_read_gbps,
        "pcie_gbps": pcie_gbps,
        "gpu_pcie_gbps": gpu_pcie_gbps,
        "target_socket_mem_read_gbps": target_socket_mem_read_gbps,
        "pcm_memory_samples": pcm_mem_samples,
        "pcm_pcie_samples": pcm_pcie_samples,
        "nvidia_dmon_rows": nvidia_dmon_rows,
        "run_seconds": run_seconds,
    }


def main() -> int:
    args = parse_args()

    files = sorted(glob.glob(os.path.join(args.root, "**/summary.txt"), recursive=True))
    print(f"Found {len(files)} summary.txt files")

    rows = []
    skipped = 0
    for f in files:
        row = parse_summary_file(f, str(args.root))
        if row:
            row["data_source"] = args.data_source
            rows.append(row)
        else:
            skipped += 1

    print(f"Parsed {len(rows)} rows, skipped {skipped}")

    if args.dry_run:
        print("DRY RUN - not writing xlsx")
        for r in rows[:3]:
            print(f"  {r['testcase']} | {r['run_label']} | {r['buffer_MiB']} MiB | "
                  f"{r['tuning_param_name']}={r['per_load_bytes']} | "
                  f"nvbw={r['nvbandwidth_gbps']} | loop={r['loop_count']} | "
                  f"mem={r['memory_gbps']} | pcie={r['pcie_gbps']}")
        return 0

    wb = openpyxl.load_workbook(args.xlsx)
    ws = wb["all_point_summary_compact"]

    assert [ws.cell(row=1, column=c).value for c in range(1, len(COLUMNS) + 1)] == COLUMNS, \
        f"Header mismatch: {[ws.cell(row=1, column=c).value for c in range(1, len(COLUMNS)+1)]}"

    next_row = ws.max_row + 1
    print(f"Appending {len(rows)} rows starting at row {next_row}")

    for i, row in enumerate(rows):
        for j, col_name in enumerate(COLUMNS):
            ws.cell(row=next_row + i, column=j + 1, value=row.get(col_name, ""))

    wb.save(args.xlsx)
    print(f"Saved {args.xlsx}")
    print(f"Total rows now: {ws.max_row}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
